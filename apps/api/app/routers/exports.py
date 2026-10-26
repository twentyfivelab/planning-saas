"""Exports (CSV, Excel, ICS, PDF) et import CSV/Excel avec aperçu. Toujours dans le tenant, filtrés."""
import csv
import io
from collections import defaultdict
from datetime import date, datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from sqlalchemy import select

from .. import models as m
from ..config import settings
from ..deps import Ctx, audit, require
from .planning import _conflicts, schedule

router = APIRouter(tags=["exports"])


def _safe(v):
    """Neutralise l'injection de formules (CSV/Excel)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _rows(ctx: Ctx, pid: str, employee_id, brand, site, date_from, date_to):
    data = schedule(pid, ctx)
    out = []
    for a in data["assignments"]:
        if employee_id and a["employee_id"] != str(employee_id):
            continue
        if site and a["site"] != site:
            continue
        if brand and brand not in a["allocations"]:
            continue
        if date_from and a["day"] < date_from.isoformat():
            continue
        if date_to and a["day"] > date_to.isoformat():
            continue
        out.append(a)
    return data, out


def _filters(employee_id: Optional[str] = Query(None), brand: Optional[str] = Query(None),
             site: Optional[str] = Query(None), date_from: Optional[date] = Query(None),
             date_to: Optional[date] = Query(None)):
    return dict(employee_id=employee_id, brand=brand, site=site, date_from=date_from, date_to=date_to)


def _job(ctx: Ctx, kind: str, f: dict):
    ctx.db.add(m.ExportJob(tenant_id=ctx.tenant_id, kind=kind, requested_by=ctx.user.id,
                           filters={k: (v.isoformat() if isinstance(v, date) else v) for k, v in f.items() if v}))
    audit(ctx, f"export.{kind}", "export", None, {k: str(v) for k, v in f.items() if v})


HEAD = ["Date", "Site", "Début", "Fin", "Présence (h)", "Pause (h)", "Effectif (h)", "Collaborateur", "Verrouillé",
        "Imputations"]


def _line(a):
    return [a["day"], a["site"], a["start"], a["end"], a["presence_minutes"] / 60, a["break_minutes"] / 60,
            a["effective_minutes"] / 60, _safe(a["employee"]), "oui" if a["locked"] else "non",
            "; ".join(f"{k}:{v / 60:g}h" for k, v in a["allocations"].items())]


@router.get("/periods/{pid}/export.csv")
def export_csv(pid: str, ctx: Ctx = Depends(require("export.run")), f: dict = Depends(_filters)):
    _, rows = _rows(ctx, pid, **f)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEAD)
    for a in rows:
        w.writerow(_line(a))
    _job(ctx, "csv", f)
    return Response(buf.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="planning.csv"'})


@router.get("/periods/{pid}/export.xlsx")
def export_xlsx(pid: str, ctx: Ctx = Depends(require("export.run")), f: dict = Depends(_filters)):
    from openpyxl import Workbook
    data, rows = _rows(ctx, pid, **f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Planning"
    ws.append(HEAD)
    for a in rows:
        ws.append(_line(a))
    ws2 = wb.create_sheet("Collaborateurs")
    ws2.append(["Collaborateur", "Jours", "Heures effectives"])
    per = defaultdict(lambda: [set(), 0])
    for a in rows:
        per[a["employee"]][0].add(a["day"])
        per[a["employee"]][1] += a["effective_minutes"]
    for k, (d, mins) in sorted(per.items()):
        ws2.append([_safe(k), len(d), mins / 60])
    ws3 = wb.create_sheet("Budgets")
    ws3.append(["Marque", "Budget (h)", "Planifié (h)", "Restant (h)", "Besoins définis"])
    planned = defaultdict(int)
    for a in rows:
        for b, mins in a["allocations"].items():
            planned[b] += mins
    for b in ctx.db.scalars(select(m.MonthlyBudget).order_by(m.MonthlyBudget.brand_code)):
        ws3.append([_safe(b.brand_code), b.target_minutes / 60, planned[b.brand_code] / 60,
                    (b.target_minutes - planned[b.brand_code]) / 60, "oui" if b.requirements_defined else "NON"])
    ws4 = wb.create_sheet("Alertes")
    ws4.append(["Type", "Gravité", "Message"])
    for u in data["uncovered"]:
        ws4.append(["Non couvert", "warning", f"{u['slot_key']} {u['start']}-{u['end']}"])
    for v in _conflicts(ctx, data["period"]["id"]):
        ws4.append(["Violation dure", "blocking", v])
    ws5 = wb.create_sheet("Hypothèses")
    ws5.append(["Hypothèse / limite"])
    for tpl in ctx.db.scalars(select(m.ShiftTemplate).where(m.ShiftTemplate.is_hypothesis.is_(True))):
        ws5.append([f"HYPOTHÈSE NON VALIDÉE : modèle {_safe(tpl.code)} {tpl.start_time:%H:%M}-{tpl.end_time:%H:%M} "
                    f"{_safe(tpl.hypothesis_note or '')}"])
    for t in ("Seuls le repos quotidien et le maximum de jours consécutifs sont exécutés (valeurs de référence non "
              "validées) ; autres règles légales suisses non exécutées : application non commercialisable en l'état.",
              "Heures budgétaires = heures de travail effectif (pause non travaillée exclue).",
              "Outil d'aide à la décision, pas un conseil juridique ; validation par un spécialiste du droit "
              "suisse requise.", "Contrats et disponibilités non renseignés = non supposés."):
        ws5.append([t])
    bio = io.BytesIO()
    wb.save(bio)
    _job(ctx, "xlsx", f)
    return Response(bio.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="planning.xlsx"'})


def _ics_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


@router.get("/periods/{pid}/export.ics")
def export_ics(pid: str, employee_id: str, ctx: Ctx = Depends(require("export.run")),
               f: dict = Depends(_filters)):
    f = {**f, "employee_id": employee_id}
    _, rows = _rows(ctx, pid, **f)
    now = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    L = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Planning SaaS//FR", "CALSCALE:GREGORIAN"]
    for a in rows:
        d = a["day"].replace("-", "")
        s, e = a["start"].replace(":", "")[:4] + "00", a["end"].replace(":", "")[:4] + "00"
        L += ["BEGIN:VEVENT", f"UID:{a['assignment_id']}@planning", f"DTSTAMP:{now}",
              f"DTSTART;TZID=Europe/Zurich:{d}T{s}", f"DTEND;TZID=Europe/Zurich:{d}T{e}",
              f"SUMMARY:{_ics_escape(a['site'])}", "END:VEVENT"]
    L.append("END:VCALENDAR")
    _job(ctx, "ics", f)
    return Response("\r\n".join(L) + "\r\n", media_type="text/calendar; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="planning.ics"'})


@router.get("/periods/{pid}/export.pdf")
def export_pdf(pid: str, ctx: Ctx = Depends(require("export.run")), f: dict = Depends(_filters)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
    data, rows = _rows(ctx, pid, **f)
    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=landscape(A4), title="Planning")
    st = getSampleStyleSheet()
    p = data["period"]
    body = [Paragraph(f"Planning {p['name']} ({p['start_date']} → {p['end_date']})", st["Title"]),
            Paragraph("Simulation exploratoire éventuelle : voir l'onglet Hypothèses de l'export Excel. "
                      "Outil d'aide à la décision, pas un conseil juridique.", st["Italic"])]
    t = Table([["Date", "Site", "Début", "Fin", "Effectif (h)", "Collaborateur"]] +
              [[a["day"], a["site"], a["start"][:5], a["end"][:5], f"{a['effective_minutes'] / 60:g}",
                a["employee"]] for a in rows], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("FONTSIZE", (0, 0), (-1, -1), 8),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.grey)]))
    body.append(t)
    doc.build(body)
    _job(ctx, "pdf", f)
    return Response(bio.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": 'attachment; filename="planning.pdf"'})


# ---------------- import ----------------
MAX_ROWS = 2000


def _parse_upload(raw: bytes, filename: str):
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        if not raw.startswith(b"PK"):
            raise HTTPException(415, "Fichier Excel invalide")
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        head = [str(c or "").strip() for c in next(it)]
        rows = [{head[i]: ("" if c is None else str(c)) for i, c in enumerate(r) if i < len(head)} for r in it]
    elif filename.lower().endswith(".csv"):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(415, "CSV non UTF-8")
        rows = list(csv.DictReader(io.StringIO(text)))
    else:
        raise HTTPException(415, "Formats acceptés : .csv, .xlsx")
    if len(rows) > MAX_ROWS:
        raise HTTPException(413, f"Maximum {MAX_ROWS} lignes")
    return rows


@router.post("/imports/employees")
async def import_employees(file: UploadFile = File(...), commit: bool = False,
                           ctx: Ctx = Depends(require("import.run"))):
    """Colonnes : code, name, [brands ex. 'Ploom GVA3:priority;Cartier:reinforce'], [weight],
    [unavailable_weekdays ex. '1;3' (0=lundi)]. Aperçu par défaut ; commit=true seulement sans erreur."""
    raw = await file.read(settings.max_upload_bytes + 1)
    if len(raw) > settings.max_upload_bytes:
        raise HTTPException(413, "Fichier trop volumineux")
    rows = _parse_upload(raw, file.filename or "")
    existing = {e.code for e in ctx.db.scalars(select(m.Employee))}
    report, valid = [], []
    seen = set()
    for i, r in enumerate(rows, start=2):
        errs = []
        code, name = (r.get("code") or "").strip(), (r.get("name") or "").strip()
        if not code or len(code) > 64:
            errs.append("code manquant ou trop long")
        if not name or len(name) > 200:
            errs.append("nom manquant ou trop long")
        if code in existing or code in seen:
            errs.append("code déjà existant")
        brands = {}
        for part in filter(None, (r.get("brands") or "").split(";")):
            b, _, lvl = part.partition(":")
            if lvl.strip() not in ("priority", "reinforce"):
                errs.append(f"niveau marque invalide: {part}")
            else:
                brands[b.strip()] = lvl.strip()
        try:
            weight = float(r.get("weight") or 1)
            if not 0 < weight <= 5:
                raise ValueError
        except ValueError:
            errs.append("weight invalide (0-5)")
            weight = 1.0
        wds = []
        for x in filter(None, (r.get("unavailable_weekdays") or "").split(";")):
            if x.strip() not in list("0123456"):
                errs.append(f"jour invalide: {x}")
            else:
                wds.append(int(x))
        seen.add(code)
        report.append({"line": i, "code": code, "errors": errs})
        if not errs:
            valid.append((code, name, brands, weight, wds))
    n_err = sum(1 for r in report if r["errors"])
    created = 0
    if commit:
        if n_err:
            raise HTTPException(422, {"message": "Corriger les erreurs avant import", "report": report})
        for code, name, brands, weight, wds in valid:
            e = m.Employee(tenant_id=ctx.tenant_id, code=code, name=name, brands=brands, allocation_weight=weight)
            ctx.db.add(e)
            ctx.db.flush()
            for wd in wds:
                ctx.db.add(m.Unavailability(tenant_id=ctx.tenant_id, employee_id=e.id, weekday=wd,
                                            reason="import"))
            created += 1
        audit(ctx, "import.employees", "employee", None, {"created": created})
    return {"rows": len(rows), "valid": len(valid), "errors": n_err, "committed": bool(commit), "created": created,
            "report": report}
