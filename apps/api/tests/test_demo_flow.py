import io

import pytest
from openpyxl import load_workbook

from app.seed import seed
from tests.conftest import make_client


@pytest.fixture(scope="module")
def demo():
    import uuid
    domain = f"{uuid.uuid4().hex[:8]}.example"
    out = seed(domain=domain)
    pw = out["password"]

    def login(email):
        c = make_client()
        r = c.post("/auth/login", json={"email": email, "password": pw})
        assert r.status_code == 200, r.text
        c.headers["X-CSRF-Token"] = r.json()["csrf_token"]
        return c
    planner, manager = login(f"planner@{domain}"), login(f"manager@{domain}")
    pid = planner.get("/periods").json()[0]["id"]
    r = planner.post(f"/periods/{pid}/runs", json={"time_limit_s": 30}, headers={"Idempotency-Key": "k-1"})
    assert r.status_code == 202, r.text
    return dict(planner=planner, manager=manager, pid=pid, run=r.json()["run_id"], login=login,
                owner=login(f"owner@{domain}"), domain=domain)


def test_run_finished_with_diagnostic(demo):
    p = demo["planner"]
    run = p.get(f"/runs/{demo['run']}").json()
    assert run["status"] == "termine" and run["result_status"] == "partiellement_couvert"
    assert run["mode"] == "simulation_exploratoire"
    diag = p.get(f"/runs/{demo['run']}/diagnostic").json()["report"]
    codes = {i["code"] for i in diag["missing_data"]}
    assert {"CONTRACT_MISSING", "BRAND_REQUIREMENTS_MISSING"} <= codes
    assert diag["weights_used"]["coverage"] == 1000
    lt = {b["brand"]: b["remaining_minutes"] for b in diag["budgets"] if not b["requirements_defined"]}
    assert lt == {"Lindt": 3840, "Toblerone": 3840}


def test_idempotent_run(demo):
    r = demo["planner"].post(f"/periods/{demo['pid']}/runs", json={"time_limit_s": 30},
                             headers={"Idempotency-Key": "k-1"})
    assert r.json()["idempotent_replay"] and r.json()["run_id"] == demo["run"]


def test_schedule_matches_solver_demo(demo):
    s = demo["planner"].get(f"/periods/{demo['pid']}/schedule").json()
    keys = [a["slot_key"].split("|")[0] for a in s["assignments"]]
    assert sum(k.startswith("GVA3") for k in keys) == 90 and sum(k.startswith("GVA9") for k in keys) == 60
    assert sum(k.startswith("CARTIER") for k in keys) == 8
    assert sum(a["effective_minutes"] for a in s["assignments"]) == 1624 * 60
    cn = [a for a in s["assignments"] if a["slot_key"].startswith("CN-")]
    assert len(cn) == 17 and all(a["allocations"] == {"Cellcosmet": 240, "Nuxe": 240} for a in cn)
    assert s["uncovered"] == []
    assert demo["planner"].get(f"/periods/{demo['pid']}/conflicts").json()["hard_violations"] == []
    assert not [a for a in s["assignments"] if a["employee"] == "Sabrina" and a["start"] in ("12:00:00", "12:30:00")]


def test_lock_and_partial_reoptimization_preserve_locked(demo):
    p = demo["planner"]
    s = p.get(f"/periods/{demo['pid']}/schedule").json()
    target = next(a for a in s["assignments"] if a["slot_key"].startswith("GVA3") and a["employee"] != "Zaki")
    assert p.post(f"/assignments/{target['assignment_id']}/lock").json()["locked"]
    assert p.put(f"/assignments/{target['assignment_id']}", json={"employee_id": target["employee_id"]}
                 ).status_code == 409
    r = p.post(f"/periods/{demo['pid']}/reoptimize", json={"time_limit_s": 30})
    assert r.status_code == 202
    after = p.get(f"/periods/{demo['pid']}/schedule").json()
    same = next(a for a in after["assignments"] if a["slot_key"] == target["slot_key"])
    assert same["employee_id"] == target["employee_id"] and same["locked"]
    before = {a["slot_key"]: a["employee"] for a in s["assignments"]}
    changed = sum(1 for a in after["assignments"] if before.get(a["slot_key"]) != a["employee"])
    assert changed < 60  # stabilité : la réoptimisation ne rebat pas tout le planning


def test_manual_move_reports_conflicts_immediately(demo):
    p = demo["planner"]
    s = p.get(f"/periods/{demo['pid']}/schedule").json()
    a = next(x for x in s["assignments"] if x["slot_key"].startswith("GVA9") and not x["locked"])
    other = next(x for x in s["assignments"] if x["day"] == a["day"] and x["employee_id"] != a["employee_id"]
                 and x["slot_key"].startswith("GVA3") and not x["locked"])
    r = p.put(f"/assignments/{a['assignment_id']}", json={"employee_id": other["employee_id"]})
    assert r.status_code == 200 and any("OVERLAP" in c or "REST" in c for c in r.json()["conflicts"])


def test_versions_compare_and_publish_refused_when_exploratory(demo):
    p = demo["planner"]
    v1 = p.post(f"/periods/{demo['pid']}/versions").json()
    assert v1["exploratory"] is True
    s = p.get(f"/periods/{demo['pid']}/schedule").json()
    a = next(x for x in s["assignments"] if x["slot_key"].startswith("GVA3") and not x["locked"])
    emps = sorted(p.get("/employees").json(), key=lambda e: e["code"])
    alt = next(e for e in emps if e["id"] != a["employee_id"])  # n'importe quel autre collaborateur : pas d'hypothèse sur un prénom
    moved = p.put(f"/assignments/{a['assignment_id']}", json={"employee_id": alt["id"]})
    assert moved.status_code == 200, moved.text
    v2 = p.post(f"/periods/{demo['pid']}/versions").json()
    cmp_ = p.get("/versions/compare", params={"a": v1["id"], "b": v2["id"]}).json()
    assert len(cmp_["changed"]) == 1
    assert p.post(f"/versions/{v2['id']}/publish").status_code == 403  # le planificateur ne publie pas
    assert demo["manager"].post(f"/versions/{v2['id']}/publish").status_code == 409  # exploratoire


def test_exports_filtered_and_scoped(demo):
    p, pid = demo["planner"], demo["pid"]
    emps = {e["code"]: e["id"] for e in p.get("/employees").json()}
    full = p.get(f"/periods/{pid}/export.csv").text.strip().splitlines()
    one = p.get(f"/periods/{pid}/export.csv", params={"employee_id": emps["jane"]}).text.strip().splitlines()
    assert len(full) > len(one) > 1 and all("Jane" in l for l in one[1:])
    cart = p.get(f"/periods/{pid}/export.csv", params={"site": "Cartier"}).text.strip().splitlines()
    assert len(cart) - 1 <= 8
    wb = load_workbook(io.BytesIO(p.get(f"/periods/{pid}/export.xlsx").content))
    assert wb.sheetnames == ["Planning", "Collaborateurs", "Budgets", "Alertes", "Hypothèses"]
    ics = p.get(f"/periods/{pid}/export.ics", params={"employee_id": emps["jane"]}).text
    assert ics.count("BEGIN:VEVENT") == len(one) - 1 and "TZID=Europe/Zurich" in ics
    assert p.get(f"/periods/{pid}/export.pdf").content.startswith(b"%PDF")
    emp_user = demo["login"](f"collaborateur@{demo['domain']}")
    assert emp_user.get(f"/periods/{pid}/export.csv").status_code == 403
    other = make_client()
    import uuid
    from tests.conftest import register
    register(other, "Autre tenant", f"{uuid.uuid4().hex}@example.com")
    assert other.get(f"/periods/{pid}/export.csv").status_code == 404


def test_cancel_and_audit(demo):
    p = demo["planner"]
    assert p.post(f"/runs/{demo['run']}/cancel").status_code == 409  # déjà terminée
    log = demo["manager"].get("/audit").json()
    actions = {a["action"] for a in log}
    assert {"optimization.requested", "optimization.finished", "assignment.lock"} <= actions
    assert "password" not in str(log).lower()


def test_no_contracts_invented_in_seed(demo):
    assert demo["planner"].get("/contracts").json() == []


def test_shift_edit_recomputes_hours_server_side(demo):
    p = demo["planner"]
    s = p.get(f"/periods/{demo['pid']}/schedule").json()
    a = next(x for x in s["assignments"] if x["slot_key"].startswith("CN-") and not x["locked"])
    r = p.patch(f"/shifts/{a['shift_id']}", json={"start_time": "09:00:00", "end_time": "19:00:00",
                                                  "break_minutes": 60})
    assert r.status_code == 200 and r.json()["presence_minutes"] == 600 and r.json()["effective_minutes"] == 540
    again = p.get(f"/periods/{demo['pid']}/schedule").json()
    cn = next(x for x in again["assignments"] if x["assignment_id"] == a["assignment_id"])
    assert sum(cn["allocations"].values()) == 540 and set(cn["allocations"]) == {"Cellcosmet", "Nuxe"}
    assert p.patch(f"/shifts/{a['shift_id']}", json={"start_time": "09:00", "end_time": "10:00",
                                                     "break_minutes": 120}).status_code == 422


def test_assumed_schedules_flagged_everywhere(demo):
    p = demo["planner"]
    tpl = {t["code"]: t for t in p.get("/shift-templates").json()}
    assert tpl["CAILLER-09"]["is_hypothesis"] and tpl["CN-09"]["is_hypothesis"]
    assert not tpl["VALMONT-08"]["is_hypothesis"] and not tpl["GVA3-06"]["is_hypothesis"]
    report = p.get(f"/runs/{demo['run']}/diagnostic").json()["report"]
    hyp = [i["message"] for i in report["missing_data"] if i["code"] == "HYPOTHESIS"]
    assert any("CAILLER-09" in h and "NON VALIDÉ" in h for h in hyp) and any("CN-09" in h for h in hyp)
    legal = next(i for i in report["missing_data"] if i["code"] == "LEGAL_REVIEW")
    assert "Non commercialisable" in legal["message"]
    wb = load_workbook(io.BytesIO(p.get(f"/periods/{demo['pid']}/export.xlsx").content))
    rows = [r[0] for r in wb["Hypothèses"].iter_rows(values_only=True)]
    assert any("HYPOTHÈSE NON VALIDÉE" in str(r) and "CAILLER-09" in str(r) for r in rows)
    assert any("non commercialisable" in str(r) for r in rows)
